# 棋子一览表（Class B roster）：65 名棋子全字段 + 来源页
import json, io, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

units = json.load(io.open('tools/units.json', encoding='utf8'))

HDR = PatternFill('solid', fgColor='1F4E79')
BAND = PatternFill('solid', fgColor='D9E1F2')
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
costColor = {1:'99DDCC', 2:'7FC9A0', 3:'B79AE8', 4:'E6B45A', 5:'FFD24A'}

wb = Workbook()
ws = wb.active
ws.title = '棋子一览'

cols = [
    ('ID', 'id', 10),
    ('名称', 'name', 12),
    ('费用', 'cost', 6),
    ('卡池张数', 'pool', 9),
    ('阵营', 'fac', 10),
    ('职业', 'job', 8),
    ('生命', 'hp', 7),
    ('攻击', 'atk', 7),
    ('射程', 'rng', 6),
    ('攻速', 'spd', 6),
    ('攻击形式', 'form', 9),
    ('伤害类型', 'dtype', 9),
    ('技能名', 'skName', 12),
    ('主动/被动', 'passive', 9),
    ('技能原型', 'arch', 10),
    ('倍率乘数', 'pow', 8),
    ('索敌偏好', 'tgt', 8),
    ('独有特效(fx)', 'fx', 26),
    ('被动参数(p)', 'p', 30),
    ('触发次数/夺盾', 'rapidN', 8),
    ('差异说明', 'txt', 34),
    ('技能描述(1★)', 'desc', 60),
]
r0 = 3
ws.cell(1, 1, '虚拟棋战 · 全棋子一览（删减用底表，2026-09-12 平衡版）').font = Font(bold=True, size=13, color='1F4E79')
for c, (t, _, w) in enumerate(cols, 1):
    cell = ws.cell(r0, c, t)
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = HDR
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border
    ws.column_dimensions[get_column_letter(c)].width = w

for i, u in enumerate(units):
    r = r0 + 1 + i
    for c, (_, k, _) in enumerate(cols, 1):
        v = u.get(k, '')
        if k in ('fx', 'p') and isinstance(v, str) and v:
            try: v = json.dumps(json.loads(v), ensure_ascii=False)
            except Exception: pass
        cell = ws.cell(r, c, v)
        cell.border = border
        cell.alignment = Alignment(horizontal='center' if k in ('cost','pool','hp','atk','rng','spd','passive','arch','pow','tgt','rapidN','form','dtype') else 'left',
                                   vertical='center', wrap_text=k in ('fx','p','txt','desc'))
        if k == 'cost':
            cell.fill = PatternFill('solid', fgColor=costColor.get(v, 'FFFFFF'))
        if i % 2 == 1 and k != 'cost':
            cell.fill = BAND

ws.freeze_panes = 'C4'
ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(cols))}{r0+len(units)}'
for rr in range(r0+1, r0+1+len(units)):
    ws.row_dimensions[rr].height = 30

# 汇总行
total = r0 + 1 + len(units)
ws.cell(total, 1, '合计').font = Font(bold=True)
ws.cell(total, 4, f'=SUM(D{r0+1}:D{total-1})').font = Font(bold=True)

# 来源页
src = wb.create_sheet('来源')
src.append(['编号','一手/二手','发布主体 · 文档或系统名','指标/字段','日期','检索于','URL'])
for c in range(1, 8):
    src.cell(1, c).font = Font(bold=True, color='FFFFFF'); src.cell(1, c).fill = HDR
src.append([1, '一手', '游戏本体 F:/demo/autochess/index.html · UNITS/SKILL_VAR 数据表',
            '名称/费用/阵营/职业/生命/攻击/射程/攻速/技能/变体/卡池张数', '2026-09-12', '2026-09-12', '—'])
src.append([])
notes = ['口径与局限：',
         '· 数据为 2026-09-12 平衡调整后的当前生效值（费用倍率、个体差异、技能系数已折算进 生命/攻击 列）。',
         '· 生命/攻击为 1★ 基础面板（未乘 HP_SCALE=2 的战斗内放大）；战斗内实际值 = 面板 ×2。',
         '· 技能描述为 1★ 文本，升星按 SKILL_STAR_M 放大。',
         '· 独有特效(fx)/被动参数(p) 为程序内部字段 JSON，供删减时定位对应 SKILL_VAR 条目。',
         '· 删减棋子时需同步删除 index.html 中 UNITS 对应行与 SKILL_VAR 对应条目，并检查其阵营/职业羁绊是否还有人凑档。']
for i, n in enumerate(notes):
    src.cell(6 + i, 1, n).font = Font(bold=(i == 0), color='666666' if i else '1F4E79')

wb.save('棋子一览_删减底表_20260912.xlsx')
print('saved, units =', len(units))
